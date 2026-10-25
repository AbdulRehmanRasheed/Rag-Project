import io
import json
import shutil
import time
import uuid
from datetime import datetime
from pathlib import Path

import streamlit as st
from docx import Document as DocxDocument
from langchain_community.vectorstores import FAISS
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_google_genai import ChatGoogleGenerativeAI, GoogleGenerativeAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

st.set_page_config(page_title="Docsight", page_icon="📚", layout="wide")

DATA_DIR = Path("chat_data")  # every chat (messages + search index) is saved here
DATA_DIR.mkdir(exist_ok=True)

# Make sure the light theme config exists (Streamlit reads it at startup, so restart once after the first run)
_cfg = Path(".streamlit/config.toml")
if not _cfg.exists():
    _cfg.parent.mkdir(exist_ok=True)
    _cfg.write_text(
        '[theme]\nbase = "light"\nprimaryColor = "#1F6F6B"\nbackgroundColor = "#FAF6EF"\n'
        'secondaryBackgroundColor = "#F1EADD"\ntextColor = "#1F2A2E"\n',
        encoding="utf-8",
    )

TYPE_STYLE = {
    "pdf": ("PDF", "#C8553D"),
    "docx": ("DOCX", "#2E6F8E"),
    "txt": ("TXT", "#B8862B"),
    "json": ("JSON", "#5B7F4A"),
}

# Google retires models fairly often. If one stops working, pick "Custom..." in the sidebar
# and type the new model name, or edit this list.
CHAT_MODELS = ["gemini-3.5-flash-lite", "gemini-3.5-flash", "gemini-3.6-flash", "gemini-3.7-flash"]
EMBED_MODEL = "models/gemini-embedding-001"

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Fraunces:opsz,wght@9..144,600;9..144,700&family=DM+Sans:wght@400;500;600&display=swap');

:root {
    --cream: #FAF6EF;
    --sand: #F1EADD;
    --ink: #1F2A2E;
    --muted: #6B7479;
    --teal: #1F6F6B;
    --teal-deep: #14403F;
    --clay: #C8553D;
    --amber: #E3A33B;
    --line: #E1D8C5;
}
:root, .stApp { color-scheme: light; }

html, body, [class*="css"], .stApp { font-family: 'DM Sans', sans-serif; color: var(--ink); }
.stApp { background: var(--cream); }
/* Hide the menu/deploy clutter but KEEP the header, because the "open sidebar" button lives in it */
#MainMenu, footer, [data-testid="stToolbar"], [data-testid="stDecoration"], [data-testid="stStatusWidget"] { visibility: hidden; }
header[data-testid="stHeader"] { background: transparent !important; }

/* The button that reopens a collapsed sidebar: make it easy to spot */
[data-testid="stExpandSidebarButton"], [data-testid="stSidebarCollapsedControl"] button,
[data-testid="collapsedControl"] button {
    visibility: visible !important; background: var(--teal) !important; border-radius: 10px !important;
    box-shadow: 0 6px 14px rgba(31,111,107,.25); transition: transform .15s ease;
}
[data-testid="stExpandSidebarButton"]:hover, [data-testid="stSidebarCollapsedControl"] button:hover,
[data-testid="collapsedControl"] button:hover { transform: scale(1.08); }
[data-testid="stExpandSidebarButton"] *, [data-testid="stSidebarCollapsedControl"] button *,
[data-testid="collapsedControl"] button * { color: #fff !important; fill: #fff !important; }
.block-container { padding-top: 1.6rem; max-width: 1100px; }

/* ---------- Animations ---------- */
@keyframes fadeUp   { from { opacity: 0; transform: translateY(14px); } to { opacity: 1; transform: none; } }
@keyframes slideIn  { from { opacity: 0; transform: translateX(-12px); } to { opacity: 1; transform: none; } }
@keyframes floatA   { 0%,100% { transform: translate(0,0) scale(1); } 50% { transform: translate(-30px,18px) scale(1.12); } }
@keyframes floatB   { 0%,100% { transform: translate(0,0) scale(1); } 50% { transform: translate(26px,-16px) scale(1.08); } }
@keyframes dots     { 0%,80%,100% { transform: scale(.5); opacity: .35; } 40% { transform: scale(1); opacity: 1; } }
@keyframes shimmer  { 0% { background-position: 0% 50%; } 100% { background-position: 200% 50%; } }

@media (prefers-reduced-motion: reduce) { * { animation: none !important; transition: none !important; } }

/* ---------- Sidebar ---------- */
section[data-testid="stSidebar"] { background: var(--sand); border-right: 1px solid var(--line); }
section[data-testid="stSidebar"] * { color: var(--ink); }
section[data-testid="stSidebar"] h2, section[data-testid="stSidebar"] h3 {
    font-family: 'Fraunces', serif; color: var(--teal-deep);
}
section[data-testid="stSidebar"] a { color: var(--teal) !important; }
.side-title { font-family: 'Fraunces', serif; font-weight: 700; color: var(--teal-deep);
              font-size: 1.15rem; margin: 1.1rem 0 .4rem 0; }

/* Labels everywhere */
[data-testid="stWidgetLabel"] p, label p, label { color: var(--ink) !important; font-weight: 500; }

/* ---------- Inputs / selects (fixes invisible text) ---------- */
div[data-baseweb="input"], div[data-baseweb="base-input"] {
    background: #fff !important; border-radius: 10px !important;
}
div[data-baseweb="input"] { border: 1px solid #D5CAB1 !important; }
div[data-baseweb="input"]:focus-within { border-color: var(--teal) !important; box-shadow: 0 0 0 3px rgba(31,111,107,.15); }
input, textarea {
    color: var(--ink) !important; -webkit-text-fill-color: var(--ink) !important;
    caret-color: var(--teal) !important; background: transparent !important;
}
input::placeholder, textarea::placeholder { color: #8A9195 !important; -webkit-text-fill-color: #8A9195 !important; }

div[data-baseweb="select"] > div {
    background: #fff !important; border: 1px solid #D5CAB1 !important; border-radius: 10px !important;
}
div[data-baseweb="select"] * { color: var(--ink) !important; -webkit-text-fill-color: var(--ink) !important; }
div[data-baseweb="select"] svg { fill: var(--teal-deep) !important; }
div[data-baseweb="popover"] ul, div[data-baseweb="popover"] [role="listbox"], div[data-baseweb="menu"] {
    background: #fff !important; border: 1px solid var(--line); border-radius: 10px;
}
div[data-baseweb="popover"] li, div[data-baseweb="popover"] [role="option"] {
    background: #fff !important; color: var(--ink) !important;
}
div[data-baseweb="popover"] li:hover, div[data-baseweb="popover"] [role="option"]:hover,
div[data-baseweb="popover"] [aria-selected="true"] { background: var(--sand) !important; }
div[data-baseweb="popover"] li * { color: var(--ink) !important; }

/* Slider */
[data-testid="stSlider"] [role="slider"] { background: var(--teal) !important; box-shadow: none !important; }
[data-testid="stSliderThumbValue"], [data-testid="stTickBarMin"], [data-testid="stTickBarMax"] {
    color: var(--teal-deep) !important; font-weight: 600;
}

/* Buttons */
.stButton > button {
    background: #fff !important; color: var(--teal-deep) !important;
    border: 1.5px solid var(--teal) !important; border-radius: 10px; font-weight: 600;
    transition: transform .15s ease, background .2s ease, box-shadow .2s ease;
}
.stButton > button:hover {
    background: var(--teal) !important; color: #fff !important;
    transform: translateY(-1px); box-shadow: 0 6px 14px rgba(31,111,107,.22);
}
.stButton > button p { color: inherit !important; }

/* Chat history buttons in sidebar */
section[data-testid="stSidebar"] .hist .stButton > button {
    border: 1px solid transparent !important; background: transparent !important;
    justify-content: flex-start; text-align: left; font-weight: 500;
    animation: slideIn .35s ease both;
}
section[data-testid="stSidebar"] .hist .stButton > button:hover { background: #fff !important; color: var(--teal-deep) !important; box-shadow: none; transform: translateX(3px); }
section[data-testid="stSidebar"] .hist.active .stButton > button {
    background: #fff !important; border-color: var(--teal) !important;
}

/* ---------- Hero ---------- */
.hero {
    position: relative; overflow: hidden;
    background: linear-gradient(135deg, var(--teal-deep) 0%, var(--teal) 100%);
    border-radius: 22px; padding: 2.4rem 2.6rem; margin-bottom: 1.4rem;
    box-shadow: 0 10px 30px rgba(20,64,63,.18);
    animation: fadeUp .6s ease both;
}
.hero::before, .hero::after { content: ""; position: absolute; border-radius: 50%; filter: blur(8px); pointer-events: none; }
.hero::before { width: 300px; height: 300px; right: -70px; top: -110px;
                background: radial-gradient(circle, rgba(227,163,59,.55), transparent 65%);
                animation: floatA 9s ease-in-out infinite; }
.hero::after  { width: 340px; height: 340px; right: 120px; bottom: -190px;
                background: radial-gradient(circle, rgba(200,85,61,.55), transparent 65%);
                animation: floatB 11s ease-in-out infinite; }
.hero > * { position: relative; z-index: 1; }
.hero h1 { font-family: 'Fraunces', serif; font-weight: 700; font-size: 2.5rem;
           margin: 0 0 .35rem 0; color: #FBF7EE !important; letter-spacing: -0.5px; }
.hero p { margin: 0; font-size: 1.05rem; color: #FBF7EE; opacity: .92; max-width: 620px; }
.hero .pills { margin-top: 1.1rem; display: flex; gap: .5rem; flex-wrap: wrap; }
.hero .pill { font-size: .75rem; font-weight: 600; letter-spacing: .6px; color: #FBF7EE;
              padding: .28rem .7rem; border-radius: 999px;
              background: rgba(255,255,255,.14); border: 1px solid rgba(255,255,255,.28);
              animation: fadeUp .6s ease both; transition: transform .2s ease, background .2s ease; }
.hero .pill:nth-child(1) { animation-delay: .25s; } .hero .pill:nth-child(2) { animation-delay: .35s; }
.hero .pill:nth-child(3) { animation-delay: .45s; } .hero .pill:nth-child(4) { animation-delay: .55s; }
.hero .pill:hover { transform: translateY(-3px); background: rgba(255,255,255,.26); }

/* ---------- Stat cards ---------- */
.stat { background: #fff; border: 1px solid var(--line); border-radius: 16px;
        padding: 1rem 1.2rem; border-top: 4px solid var(--teal);
        animation: fadeUp .55s ease both; transition: transform .2s ease, box-shadow .2s ease; }
.stat:hover { transform: translateY(-4px); box-shadow: 0 12px 24px rgba(31,42,46,.08); }
.stat.clay { border-top-color: var(--clay); animation-delay: .1s; }
.stat.amber { border-top-color: var(--amber); animation-delay: .2s; }
.stat .num { font-family: 'Fraunces', serif; font-size: 1.9rem; font-weight: 700; line-height: 1.1; color: var(--ink); }
.stat .lbl { color: var(--muted); font-size: .82rem; margin-top: .15rem; }

/* ---------- File rows ---------- */
.filerow { display: flex; align-items: center; gap: .75rem; background: #fff;
           border: 1px solid var(--line); border-radius: 12px; padding: .6rem .9rem; margin-bottom: .45rem;
           animation: slideIn .4s ease both; }
.badge { color: #fff !important; font-size: .68rem; font-weight: 700; letter-spacing: .8px;
         padding: .22rem .5rem; border-radius: 6px; min-width: 46px; text-align: center; }
.filerow .name { font-weight: 500; flex: 1; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; color: var(--ink); }
.filerow .meta { color: var(--muted); font-size: .8rem; }

/* ---------- Source cards ---------- */
.src { background: #fff; border: 1px solid var(--line); border-left: 4px solid var(--amber);
       border-radius: 10px; padding: .7rem .9rem; margin-bottom: .55rem; font-size: .88rem; color: var(--ink);
       animation: fadeUp .4s ease both; }
.src .tag { font-weight: 600; color: var(--teal-deep); font-size: .8rem; margin-bottom: .25rem; }

/* ---------- Uploader ---------- */
[data-testid="stFileUploader"] section { background: #fff !important; border: 2px dashed #CBBF9F; border-radius: 16px;
                                          transition: border-color .2s ease, background .2s ease; }
[data-testid="stFileUploader"] section:hover { border-color: var(--teal); background: #FFFDF8 !important; }
[data-testid="stFileUploader"] section * { color: var(--ink) !important; }
[data-testid="stFileUploader"] section button { background: var(--teal) !important; border: none !important; }
[data-testid="stFileUploader"] section button * { color: #fff !important; }
[data-testid="stFileUploaderFile"] { background: var(--sand) !important; border-radius: 10px; animation: slideIn .35s ease both; }
[data-testid="stFileUploaderFile"] * { color: var(--ink) !important; }
[data-testid="stFileUploaderFile"] svg { fill: var(--teal-deep) !important; }

/* ---------- Chat ---------- */
[data-testid="stChatMessage"] { background: #fff !important; border: 1px solid var(--line);
                                border-radius: 16px; padding: .9rem 1rem; animation: fadeUp .4s ease both; }
[data-testid="stChatMessage"] * { color: var(--ink); }
[data-testid="stBottom"], [data-testid="stBottom"] > div, [data-testid="stBottomBlockContainer"] {
    background: var(--cream) !important;
}
[data-testid="stChatInput"], [data-testid="stChatInput"] > div {
    background: #fff !important; border-radius: 16px !important;
}
[data-testid="stChatInput"] { border: 1.5px solid #D5CAB1; box-shadow: 0 6px 20px rgba(31,42,46,.06); transition: border-color .2s ease, box-shadow .2s ease; }
[data-testid="stChatInput"]:focus-within { border-color: var(--teal); box-shadow: 0 6px 22px rgba(31,111,107,.18); }
[data-testid="stChatInput"] textarea { color: var(--ink) !important; -webkit-text-fill-color: var(--ink) !important; }
[data-testid="stChatInput"] button { background: var(--teal) !important; border-radius: 10px; }
[data-testid="stChatInput"] button svg { fill: #fff !important; color: #fff !important; }

[data-testid="stExpander"] { background: #fff !important; border: 1px solid var(--line) !important; border-radius: 12px; }
[data-testid="stExpander"] summary * { color: var(--teal-deep) !important; font-weight: 600; }
[data-testid="stAlert"] { border-radius: 12px; animation: fadeUp .4s ease both; }
[data-testid="stAlert"] * { color: inherit; }

/* Typing indicator */
.typing { display: inline-flex; gap: 6px; padding: .35rem 0; }
.typing span { width: 9px; height: 9px; border-radius: 50%; background: var(--teal); animation: dots 1.2s infinite ease-in-out; }
.typing span:nth-child(2) { animation-delay: .15s; background: var(--amber); }
.typing span:nth-child(3) { animation-delay: .3s; background: var(--clay); }

.empty { text-align: center; color: var(--muted); padding: 2rem 1rem;
         border: 1px dashed #D5CAB1; border-radius: 16px; background: rgba(255,255,255,.55);
         animation: fadeUp .6s ease both; }
.empty b { color: var(--teal-deep); }

/* =====================================================================
   STRONG OVERRIDES: force light colours on every nested widget element
   ===================================================================== */
[data-baseweb="select"], [data-baseweb="select"] *,
[data-baseweb="input"], [data-baseweb="input"] *,
[data-baseweb="base-input"], [data-baseweb="base-input"] * {
    background-color: #FFFFFF !important;
    color: var(--ink) !important;
    -webkit-text-fill-color: var(--ink) !important;
}
[data-baseweb="select"] svg, [data-baseweb="input"] svg { fill: var(--teal-deep) !important; color: var(--teal-deep) !important; }
[data-baseweb="select"] > div, [data-baseweb="input"] { border: 1px solid #D5CAB1 !important; border-radius: 10px !important; }
[data-baseweb="select"]:hover > div, [data-baseweb="input"]:focus-within { border-color: var(--teal) !important; }

/* Dropdown list (opens in a portal at the page root) */
[data-baseweb="popover"], [data-baseweb="popover"] div, [data-baseweb="popover"] ul,
[data-baseweb="popover"] li, [data-baseweb="menu"], [data-baseweb="menu"] li {
    background-color: #FFFFFF !important; color: var(--ink) !important;
    -webkit-text-fill-color: var(--ink) !important;
}
[data-baseweb="popover"] li:hover, [data-baseweb="popover"] li:hover *,
[data-baseweb="popover"] li[aria-selected="true"], [data-baseweb="popover"] li[aria-selected="true"] * {
    background-color: var(--sand) !important;
}

/* File uploader: dropzone + the file chip inside it */
[data-testid="stFileUploader"] section *:not(button):not(svg):not(path) {
    background-color: transparent !important; color: var(--ink) !important;
    -webkit-text-fill-color: var(--ink) !important;
}
[data-testid="stFileUploader"] section [data-testid*="File"] {
    background-color: var(--sand) !important; border-radius: 10px;
}
[data-testid="stFileUploader"] section button, [data-testid="stFileUploader"] section button * {
    background-color: var(--teal) !important; color: #fff !important; -webkit-text-fill-color: #fff !important;
}
[data-testid="stFileUploader"] svg { fill: var(--teal-deep) !important; }
[data-testid="stFileUploader"] section button svg { fill: #fff !important; }

/* Slider fill in teal */
[data-testid="stSlider"] div[data-baseweb="slider"] div { background-color: transparent; }
[data-testid="stSlider"] [role="slider"] { background-color: var(--teal) !important; }

/* =====================================================================
   Chat history cards in the sidebar
   ===================================================================== */
[class*="st-key-chat_"] {
    background: #fff; border: 1px solid var(--line); border-radius: 12px;
    padding: .3rem .4rem .1rem .4rem; margin-bottom: .5rem;
    animation: slideIn .35s ease both; transition: transform .15s ease, box-shadow .2s ease;
}
[class*="st-key-chat_"]:hover { transform: translateX(3px); box-shadow: 0 6px 14px rgba(31,42,46,.08); }
[class*="st-key-chat_on_"] { border: 1.5px solid var(--teal); background: #F4FAF8; }

[class*="st-key-open_"] button {
    background: transparent !important; border: none !important; box-shadow: none !important;
    justify-content: flex-start !important; text-align: left; padding: .25rem .4rem; min-height: 0;
}
[class*="st-key-open_"] button:hover { background: var(--sand) !important; transform: none; }
[class*="st-key-open_"] button p {
    white-space: nowrap; overflow: hidden; text-overflow: ellipsis; max-width: 150px;
    color: var(--ink) !important; font-weight: 600; font-size: .9rem;
}
[class*="st-key-del_"] button {
    background: transparent !important; border: none !important; box-shadow: none !important;
    color: var(--clay) !important; padding: .1rem .3rem; min-height: 0;
}
[class*="st-key-del_"] button:hover { background: #FBE9E4 !important; color: var(--clay) !important; transform: none; }
[class*="st-key-chat_"] [data-testid="stCaptionContainer"] { padding-left: .45rem; margin-top: -.4rem; }
[class*="st-key-chat_"] [data-testid="stCaptionContainer"] p { color: var(--muted) !important; font-size: .75rem; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

TYPING_HTML = '<div class="typing"><span></span><span></span><span></span></div>'


# ---------------------------------------------------------------
# Saved chats on disk
# ---------------------------------------------------------------
def chat_path(cid: str) -> Path:
    return DATA_DIR / cid


def save_chat(meta: dict):
    meta["updated"] = datetime.now().isoformat(timespec="seconds")
    p = chat_path(meta["id"])
    p.mkdir(parents=True, exist_ok=True)
    (p / "chat.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def read_chat(cid: str) -> dict | None:
    f = chat_path(cid) / "chat.json"
    if not f.exists():
        return None
    try:
        return json.loads(f.read_text(encoding="utf-8"))
    except Exception:
        return None


def list_chats() -> list[dict]:
    chats = []
    for d in DATA_DIR.iterdir():
        if d.is_dir():
            m = read_chat(d.name)
            if m:
                chats.append(m)
    return sorted(chats, key=lambda m: m.get("updated", ""), reverse=True)


# ---------------------------------------------------------------
# Session state + callbacks
# ---------------------------------------------------------------
for k, v in {
    "chat_id": None, "messages": [], "vectorstore": None, "files_meta": [],
    "stats": (0, 0, 0), "sig": None, "files_sig": None, "uploader_key": 0,
    "pending_load": None, "skipped": [],
}.items():
    st.session_state.setdefault(k, v)


def reset_chat():
    st.session_state.update(
        chat_id=None, messages=[], vectorstore=None, files_meta=[], stats=(0, 0, 0),
        sig=None, files_sig=None, skipped=[], pending_load=None,
        uploader_key=st.session_state.uploader_key + 1,
    )


def request_load(cid: str):
    st.session_state.pending_load = cid


def delete_chat(cid: str):
    shutil.rmtree(chat_path(cid), ignore_errors=True)
    if st.session_state.chat_id == cid:
        reset_chat()


# ---------------------------------------------------------------
# Sidebar
# ---------------------------------------------------------------
with st.sidebar:
    st.subheader("Settings")
    api_key = st.text_input("Gemini API key", type="password", placeholder="Paste your Gemini API key")
    st.caption("Get a free key at [aistudio.google.com/apikey](https://aistudio.google.com/apikey)")

    choice = st.selectbox("Chat model", CHAT_MODELS + ["Custom..."])
    if choice == "Custom...":
        model_name = st.text_input("Model name", placeholder="e.g. gemini-3.5-flash")
    else:
        model_name = choice

    # Chats come first so they're always visible without scrolling
    st.markdown('<div class="side-title">Your chats</div>', unsafe_allow_html=True)
    st.button("＋  New chat", key="new_chat", use_container_width=True, on_click=reset_chat)

    history = list_chats()
    if not history:
        st.caption("Upload a file and your chat will be saved here automatically.")
    for m in history:
        cid = m["id"]
        active = cid == st.session_state.chat_id
        n_qa = len(m.get("messages", [])) // 2
        updated = m.get("updated", "")[:16].replace("T", " ")
        with st.container(key=f"chat_{'on' if active else 'off'}_{cid}"):
            c1, c2 = st.columns([6, 1], vertical_alignment="center")
            c1.button(
                m.get("title", "Untitled"), key=f"open_{cid}", use_container_width=True,
                on_click=request_load, args=(cid,), help=m.get("title", ""),
            )
            c2.button("✕", key=f"del_{cid}", on_click=delete_chat, args=(cid,), help="Delete this chat")
            st.caption(f"{len(m.get('files', []))} file(s) · {n_qa} question(s) · {updated}")

    with st.expander("Retrieval settings"):
        chunk_size = st.slider("Chunk size", 500, 2000, 1000, step=100)
        chunk_overlap = st.slider("Chunk overlap", 0, 400, 150, step=50)
        top_k = st.slider("Chunks per answer (k)", 2, 10, 4)

# ---------------------------------------------------------------
# Hero
# ---------------------------------------------------------------
st.markdown(
    """
    <div class="hero">
        <h1>Docsight</h1>
        <p>Drop in your documents and ask questions in plain language.
        Every answer is pulled only from your files, with the source shown.</p>
        <div class="pills">
            <span class="pill">PDF</span><span class="pill">DOCX</span>
            <span class="pill">TXT</span><span class="pill">JSON</span>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

if not api_key:
    st.info("Add your Gemini API key in the sidebar to get started.")
    st.stop()
if not model_name:
    st.info("Type a model name in the sidebar.")
    st.stop()


def get_embeddings():
    return GoogleGenerativeAIEmbeddings(model=EMBED_MODEL, google_api_key=api_key)


# ---------------------------------------------------------------
# Open a saved chat (done before the uploader is drawn)
# ---------------------------------------------------------------
if st.session_state.pending_load:
    cid = st.session_state.pending_load
    st.session_state.pending_load = None
    meta = read_chat(cid)
    if meta:
        try:
            store = FAISS.load_local(
                str(chat_path(cid) / "index"), get_embeddings(),
                allow_dangerous_deserialization=True,  # our own saved index
            )
            st.session_state.update(
                chat_id=cid, messages=meta.get("messages", []), vectorstore=store,
                files_meta=meta.get("files", []), stats=tuple(meta.get("stats", (0, 0, 0))),
                sig=None, files_sig=None, skipped=[],
                uploader_key=st.session_state.uploader_key + 1,
            )
        except Exception as e:
            st.error(f"Couldn't open that chat: {e}")


# ---------------------------------------------------------------
# Loading files -> LangChain Documents
# ---------------------------------------------------------------
def decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def flatten_json(obj, prefix=""):
    lines = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            lines += flatten_json(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            lines += flatten_json(v, f"{prefix}[{i}]")
    else:
        lines.append(f"{prefix}: {obj}")
    return lines


def load_file(name: str, raw: bytes) -> list[Document]:
    ext = name.rsplit(".", 1)[-1].lower()
    docs: list[Document] = []

    if ext == "pdf":
        reader = PdfReader(io.BytesIO(raw))
        for i, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            if text:
                docs.append(Document(page_content=text,
                                     metadata={"source": name, "label": f"{name}, p. {i + 1}"}))
    elif ext == "docx":
        d = DocxDocument(io.BytesIO(raw))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        if parts:
            docs.append(Document(page_content="\n".join(parts), metadata={"source": name, "label": name}))
    elif ext == "txt":
        text = decode_text(raw).strip()
        if text:
            docs.append(Document(page_content=text, metadata={"source": name, "label": name}))
    elif ext == "json":
        text = "\n".join(flatten_json(json.loads(decode_text(raw))))
        if text.strip():
            docs.append(Document(page_content=text, metadata={"source": name, "label": name}))
    return docs


def embed_in_batches(chunks: list[Document], embeddings, batch_size: int = 40) -> FAISS:
    """Small batches + retry, because the free Gemini tier rate-limits embeddings."""
    store = None
    bar = st.progress(0.0, text="Embedding chunks...")
    for start in range(0, len(chunks), batch_size):
        batch = chunks[start:start + batch_size]
        for attempt in range(4):
            try:
                if store is None:
                    store = FAISS.from_documents(batch, embeddings)
                else:
                    store.add_documents(batch)
                break
            except Exception:
                if attempt == 3:
                    raise
                time.sleep(6 * (attempt + 1))
        bar.progress(min((start + batch_size) / len(chunks), 1.0), text="Embedding chunks...")
    bar.empty()
    return store


def build_vectorstore(files, size, overlap):
    all_docs, skipped = [], []
    for name, raw in files:
        try:
            loaded = load_file(name, raw)
        except Exception as e:
            skipped.append(f"{name} ({e.__class__.__name__})")
            continue
        if not loaded:
            skipped.append(f"{name} (no readable text)")
        all_docs += loaded

    if not all_docs:
        return None, 0, 0, skipped

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size, chunk_overlap=overlap, separators=["\n\n", "\n", ". ", " ", ""]
    )
    chunks = splitter.split_documents(all_docs)
    store = embed_in_batches(chunks, get_embeddings())
    return store, len(chunks), sum(len(d.page_content) for d in all_docs), skipped


# ---------------------------------------------------------------
# Upload + indexing
# ---------------------------------------------------------------
uploaded = st.file_uploader(
    "Upload your documents", type=["pdf", "docx", "txt", "json"],
    accept_multiple_files=True, key=f"uploader_{st.session_state.uploader_key}",
)

if uploaded:
    files_sig = tuple((f.name, f.size) for f in uploaded)
    sig = (files_sig, chunk_size, chunk_overlap)
    if st.session_state.sig != sig:
        with st.spinner("Reading your files..."):
            try:
                store, n_chunks, n_chars, skipped = build_vectorstore(
                    [(f.name, f.getvalue()) for f in uploaded], chunk_size, chunk_overlap
                )
            except Exception as e:
                st.error(f"Indexing failed: {e}")
                st.stop()
        st.session_state.sig = sig
        st.session_state.skipped = skipped

        if store is None:
            st.session_state.vectorstore = None
            st.error("None of those files had readable text.")
        else:
            # Same files with new slider values -> keep the same saved chat, otherwise start a new one
            same_files = st.session_state.files_sig == files_sig and st.session_state.chat_id
            cid = st.session_state.chat_id if same_files else uuid.uuid4().hex[:12]
            files_meta = [{"name": f.name, "size": f.size} for f in uploaded]
            title = files_meta[0]["name"] + (f" +{len(files_meta) - 1}" if len(files_meta) > 1 else "")

            chat_path(cid).mkdir(parents=True, exist_ok=True)
            store.save_local(str(chat_path(cid) / "index"))

            st.session_state.update(
                chat_id=cid, messages=[], vectorstore=store, files_meta=files_meta,
                stats=(len(uploaded), n_chunks, n_chars), files_sig=files_sig,
            )
            save_chat({
                "id": cid, "title": title, "created": datetime.now().isoformat(timespec="seconds"),
                "files": files_meta, "stats": [len(uploaded), n_chunks, n_chars], "messages": [],
            })
            st.rerun()  # refresh sidebar list

for s in st.session_state.skipped:
    st.warning(f"Skipped {s}")

# ---------------------------------------------------------------
# Library overview
# ---------------------------------------------------------------
if st.session_state.vectorstore is not None:
    n_files, n_chunks, n_chars = st.session_state.stats
    c1, c2, c3 = st.columns(3)
    c1.markdown(f'<div class="stat"><div class="num">{n_files}</div><div class="lbl">Files indexed</div></div>', unsafe_allow_html=True)
    c2.markdown(f'<div class="stat clay"><div class="num">{n_chunks}</div><div class="lbl">Searchable chunks</div></div>', unsafe_allow_html=True)
    c3.markdown(f'<div class="stat amber"><div class="num">{n_chars:,}</div><div class="lbl">Characters read</div></div>', unsafe_allow_html=True)

    with st.expander("Your library", expanded=False):
        for f in st.session_state.files_meta:
            ext = f["name"].rsplit(".", 1)[-1].lower()
            label, color = TYPE_STYLE.get(ext, ("FILE", "#666"))
            st.markdown(
                f'<div class="filerow"><span class="badge" style="background:{color}">{label}</span>'
                f'<span class="name">{f["name"]}</span><span class="meta">{f["size"] / 1024:.1f} KB</span></div>',
                unsafe_allow_html=True,
            )

# ---------------------------------------------------------------
# RAG pieces
# ---------------------------------------------------------------
RAG_PROMPT = ChatPromptTemplate.from_messages([
    ("system",
     "You answer questions strictly using the context taken from the user's documents.\n"
     "Rules:\n"
     "1. Use ONLY the context below. Do not use outside knowledge.\n"
     "2. If the answer isn't in the context, say: \"I couldn't find that in the documents.\"\n"
     "3. Cite where you found it using the bracketed source labels, e.g. (report.pdf, p. 3).\n"
     "4. If several files are relevant, combine them and cite each one.\n\n"
     "Context:\n{context}"),
    ("human", "{question}"),
])


def format_docs(docs: list[Document]) -> str:
    return "\n\n".join(f"[{d.metadata['label']}]\n{d.page_content}" for d in docs)


def chunk_text(chunk) -> str:
    """Newer Gemini models can return a list of content blocks instead of a plain string."""
    if isinstance(chunk, str):
        return chunk
    if isinstance(chunk, list):
        return "".join(b.get("text", "") if isinstance(b, dict) else str(b) for b in chunk)
    return str(chunk)


def render_sources(sources: list[dict]):
    with st.expander(f"Sources ({len(sources)} chunks)"):
        for s in sources:
            st.markdown(f'<div class="src"><div class="tag">{s["label"]}</div>{s["text"]}</div>',
                        unsafe_allow_html=True)


# ---------------------------------------------------------------
# Chat
# ---------------------------------------------------------------
if st.session_state.vectorstore is not None:
    for msg in st.session_state.messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])
            if msg.get("sources"):
                render_sources(msg["sources"])

    question = st.chat_input("Ask something about your documents...")
    if question:
        st.session_state.messages.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        retriever = st.session_state.vectorstore.as_retriever(search_kwargs={"k": top_k})
        llm = ChatGoogleGenerativeAI(model=model_name, temperature=0, google_api_key=api_key)
        chain = RAG_PROMPT | llm

        with st.chat_message("assistant"):
            slot = st.empty()
            slot.markdown(TYPING_HTML, unsafe_allow_html=True)
            answer = ""
            try:
                docs = retriever.invoke(question)
                for chunk in chain.stream({"context": format_docs(docs), "question": question}):
                    answer += chunk_text(chunk.content)
                    slot.markdown(answer + " ▌")
                slot.markdown(answer)
            except Exception as e:
                slot.empty()
                st.session_state.messages.pop()  # drop the unanswered question
                st.error(f"Something went wrong: {e}")
                st.stop()

            sources = [{"label": d.metadata["label"],
                        "text": d.page_content[:400].replace("\n", " ")} for d in docs]
            render_sources(sources)

        st.session_state.messages.append({"role": "assistant", "content": answer, "sources": sources})

        # Save to disk so the chat is still there tomorrow
        meta = read_chat(st.session_state.chat_id)
        if meta:
            if not any(m["role"] == "user" for m in meta.get("messages", [])):
                meta["title"] = question[:48] + ("…" if len(question) > 48 else "")
            meta["messages"] = st.session_state.messages
            save_chat(meta)
            if len(st.session_state.messages) == 2:
                st.rerun()  # refresh sidebar title after the first answer
elif not uploaded:
    st.markdown(
        '<div class="empty"><b>Nothing here yet.</b><br>Upload a few files above, or open one of your saved chats from the sidebar.</div>',
        unsafe_allow_html=True,
    )
